from __future__ import annotations

from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import (
    WD_ALIGN_PARAGRAPH,
    WD_BREAK,
    WD_LINE_SPACING,
    WD_TAB_ALIGNMENT,
)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables"
OUT.mkdir(exist_ok=True)

GUIDE_PATH = OUT / "FFPMU_Country_Leader_Pilot_Training_Guide.docx"
COMMUNITY_PATH = OUT / "FFPMU_European_Community_Announcement.docx"
SCREENSHOT_DIR = ROOT / "docs" / "ffpmu_guide_screenshots"

TODAY = date(2026, 8, 6)
DATE_LABEL = "6 August 2026"
APP_URL = "https://ffpmupt-402e1.web.app/"
GITHUB_URL = "https://github.com/OsmanBrito/ffpmupt"
FAMILY_PLEDGE_URL = "https://familypledge.app/"
FAMILY_PLEDGE_IOS = "https://apps.apple.com/us/app/family-pledge/id6789864172"
FAMILY_PLEDGE_ANDROID = (
    "https://play.google.com/store/apps/details?id="
    "org.familyfederation.family_pledge&hl=en"
)

# compact_reference_guide tokens, with one named FFPMU palette override.
GREEN = "245C52"
DARK_GREEN = "193C37"
PALE_GREEN = "E7F0ED"
CREAM = "F7F5EF"
GOLD = "7A5428"
INK = "202826"
MUTED = "5B6864"
GRID = "CDD9D5"
LIGHT_GRID = "E8E1D5"
WHITE = "FFFFFF"
ERROR = "9B1C1C"


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_run_font(run, size=None, bold=None, italic=None, color=INK, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.color.rgb = rgb(color)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    return run


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=GRID, size=6):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths_dxa: list[int], indent_dxa=120):
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[min(index, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def add_hyperlink(paragraph, text, url, color=GREEN):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Calibri")
    r_fonts.set(qn("w:hAnsi"), "Calibri")
    r_pr.append(r_fonts)
    color_el = OxmlElement("w:color")
    color_el.set(qn("w:val"), color)
    r_pr.append(color_el)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_field(paragraph, instruction: str):
    def field_run(child):
        run = OxmlElement("w:r")
        r_pr = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), "Calibri")
        fonts.set(qn("w:hAnsi"), "Calibri")
        r_pr.append(fonts)
        color = OxmlElement("w:color")
        color.set(qn("w:val"), MUTED)
        r_pr.append(color)
        size = OxmlElement("w:sz")
        size.set(qn("w:val"), "18")
        r_pr.append(size)
        run.append(r_pr)
        run.append(child)
        paragraph._p.append(run)

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    field_run(begin)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    field_run(instr)
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    field_run(separate)
    text = OxmlElement("w:t")
    text.text = "1"
    field_run(text)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    field_run(end)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_page_number_footer(section, label: str):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    set_run_font(p.add_run(f"{label}  |  "), size=9, color=MUTED)
    add_field(p, "PAGE")
    set_run_font(p.add_run(" / "), size=9, color=MUTED)
    add_field(p, "NUMPAGES")


def add_running_header(section, left: str, right: str):
    header = section.header
    header.is_linked_to_previous = False
    p = header.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.tab_stops.add_tab_stop(
        Inches(6.5), WD_TAB_ALIGNMENT.RIGHT
    )
    set_run_font(p.add_run(left), size=9, bold=True, color=GREEN)
    p.add_run("\t")
    set_run_font(p.add_run(right), size=9, color=MUTED)


def configure_section_geometry(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.86)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(0.82)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.42)
    section.footer_distance = Inches(0.42)


def configure_styles(doc: Document):
    doc.settings.odd_and_even_pages_header_footer = False
    section = doc.sections[0]
    configure_section_geometry(section)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_specs = {
        "Heading 1": (16, GREEN, 18, 10),
        "Heading 2": (13, GREEN, 14, 7),
        "Heading 3": (12, DARK_GREEN, 10, 5),
    }
    for name, (size, color, before, after) in heading_specs.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.font.color.rgb = rgb(INK)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    if "Guide Kicker" not in styles:
        kicker = styles.add_style("Guide Kicker", WD_STYLE_TYPE.PARAGRAPH)
    else:
        kicker = styles["Guide Kicker"]
    kicker.base_style = normal
    kicker.font.name = "Calibri"
    kicker.font.size = Pt(10)
    kicker.font.bold = True
    kicker.font.color.rgb = rgb(GOLD)
    kicker.paragraph_format.space_after = Pt(2)
    kicker.paragraph_format.keep_with_next = True

    if "Guide Title" not in styles:
        title = styles.add_style("Guide Title", WD_STYLE_TYPE.PARAGRAPH)
    else:
        title = styles["Guide Title"]
    title.base_style = normal
    title.font.name = "Calibri"
    title.font.size = Pt(29)
    title.font.bold = True
    title.font.color.rgb = rgb(DARK_GREEN)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)
    title.paragraph_format.keep_with_next = True

    if "Guide Subtitle" not in styles:
        subtitle = styles.add_style("Guide Subtitle", WD_STYLE_TYPE.PARAGRAPH)
    else:
        subtitle = styles["Guide Subtitle"]
    subtitle.base_style = normal
    subtitle.font.name = "Calibri"
    subtitle.font.size = Pt(13.5)
    subtitle.font.color.rgb = rgb(MUTED)
    subtitle.paragraph_format.space_after = Pt(18)
    subtitle.paragraph_format.keep_with_next = True

    if "Callout" not in styles:
        callout = styles.add_style("Callout", WD_STYLE_TYPE.PARAGRAPH)
    else:
        callout = styles["Callout"]
    callout.base_style = normal
    callout.font.name = "Calibri"
    callout.font.size = Pt(10.5)
    callout.font.color.rgb = rgb(DARK_GREEN)
    callout.paragraph_format.left_indent = Inches(0.16)
    callout.paragraph_format.right_indent = Inches(0.16)
    callout.paragraph_format.space_before = Pt(8)
    callout.paragraph_format.space_after = Pt(10)
    callout.paragraph_format.line_spacing = 1.2


def add_callout(doc, label: str, text: str, fill=PALE_GREEN, border=GREEN):
    p = doc.add_paragraph(style="Callout")
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    p_bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:color"), border)
    left.set(qn("w:space"), "8")
    p_bdr.append(left)
    p_pr.append(p_bdr)
    set_run_font(p.add_run(f"{label}: "), size=10.5, bold=True, color=DARK_GREEN)
    set_run_font(p.add_run(text), size=10.5, color=DARK_GREEN)
    return p


def add_para(doc, text="", *, bold=False, italic=False, color=INK, size=11,
             align=None, before=0, after=6, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    if align is not None:
        p.alignment = align
    if keep:
        p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(text), size=size, bold=bold, italic=italic, color=color)
    return p


def add_bullet(doc, text: str, bold_prefix: str | None = None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_prefix and text.startswith(bold_prefix):
        set_run_font(p.add_run(bold_prefix), bold=True, color=DARK_GREEN)
        set_run_font(p.add_run(text[len(bold_prefix):]))
    else:
        set_run_font(p.add_run(text))
    return p


def add_number(doc, text: str, bold_prefix: str | None = None):
    previous_is_number = bool(
        doc.paragraphs and doc.paragraphs[-1].style.name == "List Number"
    )
    if not previous_is_number:
        numbering = doc.part.numbering_part.element
        style_num_id = int(
            doc.styles["List Number"]._element.pPr.numPr.numId.val
        )
        source_num = next(
            node
            for node in numbering.findall(qn("w:num"))
            if int(node.get(qn("w:numId"))) == style_num_id
        )
        abstract_id = int(source_num.find(qn("w:abstractNumId")).get(qn("w:val")))
        existing_ids = [
            int(node.get(qn("w:numId")))
            for node in numbering.findall(qn("w:num"))
        ]
        current_num_id = max(existing_ids, default=0) + 1
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(current_num_id))
        abstract = OxmlElement("w:abstractNumId")
        abstract.set(qn("w:val"), str(abstract_id))
        num.append(abstract)
        override = OxmlElement("w:lvlOverride")
        override.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:startOverride")
        start.set(qn("w:val"), "1")
        override.append(start)
        num.append(override)
        numbering.append(num)
        doc._ffpmu_current_num_id = current_num_id
    p = doc.add_paragraph(style="List Number")
    p_pr = p._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.insert(0, num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), str(doc._ffpmu_current_num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id)
    if bold_prefix and text.startswith(bold_prefix):
        set_run_font(p.add_run(bold_prefix), bold=True, color=DARK_GREEN)
        set_run_font(p.add_run(text[len(bold_prefix):]))
    else:
        set_run_font(p.add_run(text))
    return p


def add_table(doc, headers: list[str], rows: list[list[str]], widths: list[int],
              *, font_size=9.5, header_fill=PALE_GREEN):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths, indent_dxa=120)
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        shade_cell(cell, header_fill)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.12
        set_run_font(p.add_run(header), size=font_size, bold=True, color=DARK_GREEN)
    for row_data in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_data):
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.12
            set_run_font(p.add_run(value), size=font_size, color=INK)
    set_table_geometry(table, widths, indent_dxa=120)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_link_line(doc, label: str, url: str, note: str | None = None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.18
    set_run_font(p.add_run(f"{label}: "), bold=True, color=DARK_GREEN)
    add_hyperlink(p, url, url)
    if note:
        set_run_font(p.add_run(f" — {note}"), italic=True, color=MUTED, size=10)
    return p


def add_title_block(doc, kicker: str, title: str, subtitle: str,
                    metadata: list[tuple[str, str]], title_size=29):
    doc.add_paragraph(kicker.upper(), style="Guide Kicker")
    title_p = doc.add_paragraph(title, style="Guide Title")
    for run in title_p.runs:
        set_run_font(run, size=title_size, bold=True, color=DARK_GREEN)
    doc.add_paragraph(subtitle, style="Guide Subtitle")
    rows = [[metadata[0][0], metadata[0][1], metadata[1][0], metadata[1][1]],
            [metadata[2][0], metadata[2][1], metadata[3][0], metadata[3][1]]]
    table = doc.add_table(rows=2, cols=4)
    set_table_geometry(table, [1400, 3280, 1400, 3280], indent_dxa=120)
    set_table_borders(table, color=LIGHT_GRID)
    set_repeat_table_header(table.rows[0])
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            if c_idx % 2 == 0:
                shade_cell(cell, PALE_GREEN)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            set_run_font(
                p.add_run(value),
                size=9.5,
                bold=c_idx % 2 == 0,
                color=DARK_GREEN if c_idx % 2 == 0 else INK,
            )
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    p.paragraph_format.space_after = Pt(0)


def add_screenshot(doc, filename: str, caption: str, note: str):
    image_path = SCREENSHOT_DIR / filename
    if not image_path.exists():
        raise FileNotFoundError(f"Guide screenshot not found: {image_path}")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(3)
    picture = p.add_run().add_picture(str(image_path), width=Inches(6.15))
    picture._inline.docPr.set("title", caption)
    picture._inline.docPr.set("descr", f"{caption}. {note}")
    caption_p = doc.add_paragraph()
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_p.paragraph_format.space_after = Pt(2)
    set_run_font(caption_p.add_run(caption), size=9.5, bold=True, color=DARK_GREEN)
    note_p = doc.add_paragraph()
    note_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note_p.paragraph_format.space_after = Pt(9)
    set_run_font(note_p.add_run(note), size=9, italic=True, color=MUTED)


def build_guide():
    doc = Document()
    configure_styles(doc)
    add_running_header(doc.sections[0], "FFPMU | Country leaders", "Practical guide")
    add_page_number_footer(doc.sections[0], "FFPMU Country Leader Guide")

    add_title_block(
        doc,
        "Practical guide for country leaders",
        "FFPMU Country\nUser & Admin Guide",
        "A friendly, step-by-step guide to setting up your country and using FFPMU together.",
        [
            ("Audience", "National leaders and country administrators"),
            ("Edition", f"Pilot · {DATE_LABEL}"),
            ("Current channel", "Web-first pilot"),
            ("First active country", "Portugal"),
        ],
    )
    add_callout(
        doc,
        "How to use this guide",
        "Read it once, then keep it open while you explore the app. Start with the public pages, set up your country content, and try Sunday Mode on a real Sunday whenever it feels helpful. Use it at your own pace.",
    )
    add_link_line(doc, "Pilot app", APP_URL, "open in a modern browser")
    add_link_line(doc, "Source repository", GITHUB_URL, "access may need to be granted")

    doc.add_heading("1. What FFPMU is", level=1)
    add_para(
        doc,
        "FFPMU is a web-based Sunday service guide and country content hub. Members choose a country and receive that country's language and public content. The same shared application can support different national communities without mixing their administrative data.",
    )
    for item in [
        "Sunday Mode with a configurable order of service, completion tracking, full-screen presentation, and a short local use summary.",
        "Songs with lyrics, categories, optional audio, verse timing, video links, and offline audio preparation.",
        "Family Pledge in the country's main language, Korean, and English.",
        "Yearly motto, country-specific offerings/tithes, and weekly YouTube/Vimeo videos.",
        "A shared European Holy Grounds directory with country filters, search, visit information, and Google Maps links.",
    ]:
        add_bullet(doc, item)
    add_callout(
        doc,
        "Current scope",
        "The public pilot is currently a web application. Native Android and iOS releases have not been published. Videos still require internet; previously synchronized content and prepared audio support offline use.",
        fill=CREAM,
        border=GOLD,
    )

    doc.add_heading("2. First screen: choose your country", level=1)
    add_para(
        doc,
        "On the first visit, choose the country you are setting up or visiting. Only enabled countries appear here. After choosing, the address includes the country code, for example /pt, /de, or /fr.",
    )
    add_screenshot(
        doc,
        "01-country-selection.png",
        "What a member sees when choosing a country",
        "The sample countries in this image are for illustration; your app will show the countries that are currently enabled.",
    )
    add_bullet(doc, "Use the globe icon in the top bar to change country later.")
    add_bullet(doc, "Use the translate icon to change the interface language without changing the country content.")

    doc.add_heading("3. The public experience", level=1)
    add_para(doc, "Once a country is selected, members can use these pages without an admin login:")
    for item in [
        "Sunday Mode — choose the order for the local service, prepare audio, and open each part when needed.",
        "Songs — search by title, page, or lyric; view lyrics and play any available audio.",
        "Family Pledge — switch between the country language, Korean, and English, then move through the eight points.",
        "Motto, offerings, weekly videos, and European Holy Grounds — country content and shared European information in one place.",
    ]:
        add_bullet(doc, item)
    add_screenshot(
        doc,
        "02-motto.png",
        "Example of a public content screen: the yearly motto",
        "Screens adapt to the selected country and interface language. The icon may vary slightly by browser or device.",
    )

    doc.add_heading("4. Roles and country boundaries", level=1)
    add_table(
        doc,
        ["Role", "What this person can do", "Country boundary"],
        [
            ["Member / public user", "Use Sunday Mode and view public songs, Pledge, motto, offerings, videos, and Holy Grounds.", "No editing access."],
            ["Country administrator", "Manage settings, readiness, Pledge, songs, payments, videos, and Holy Grounds.", "Only the exact country codes assigned to the account."],
            ["Superadministrator", "Onboard countries, local churches, and country admins; open any country's admin area.", "Global access; keep this role limited to a few coordinators."],
        ],
        [1750, 4700, 2910],
        font_size=9.2,
    )
    add_para(
        doc,
        "Country content is independent. Editing or disabling a song, Pledge language, payment method, or Holy Ground in one country does not change another country.",
        bold=True,
        color=DARK_GREEN,
    )
    add_bullet(doc, "The operational CRM stores country, church, administrator, and setup-readiness information.")
    add_bullet(doc, "This version does not store member profiles, attendance, or pastoral data.")
    add_bullet(doc, "Sunday Mode plans and service reports stay in the local browser/device storage.")

    doc.add_heading("5. What to collect before the first setup", level=1)
    add_para(doc, "It is easiest to gather these details before the country is created. You can also add them gradually as your team prepares them:")
    for item in [
        "Country name and lowercase country code (for example, de, fr, es, or it).",
        "Main interface language and IANA timezone (for example, Europe/Berlin).",
        "National administrator's name and individual email address.",
        "At least one local church: name, city, address, timezone, responsible person, and contact email.",
        "Approved Family Pledge translation with exactly eight complete points.",
        "Initial songs and lyrics, preferably using the provided XLSX template or one PPTX per song.",
        "Approved offering/tithing details and any QR payload or payment URL.",
        "Current YouTube and Vimeo weekly-video links.",
        "Optional Holy Ground information: name, city, address, image, story, visit instructions, contact, and coordinates.",
    ]:
        add_bullet(doc, item)

    doc.add_heading("Supported interface languages", level=2)
    add_table(
        doc,
        ["Language", "Code", "Language", "Code"],
        [
            ["Portuguese", "pt", "Portuguese (Brazil)", "pt-br"],
            ["English", "en", "Korean", "ko"],
            ["Spanish", "es", "German", "de"],
            ["Italian", "it", "French", "fr"],
        ],
        [2700, 900, 2700, 3060],
    )
    add_callout(
        doc,
        "Language distinction",
        "The country's main language sets the default interface. Users may still change the interface language with the translate icon. Content such as songs and Holy Ground stories is entered by administrators and is not translated automatically.",
    )

    doc.add_heading("6. Administrator invitation and first sign-in", level=1)
    doc.add_heading("For the superadministrator", level=2)
    steps = [
        ("Open the correct country.", "Go to Administration > Operational CRM, then open the country's card."),
        ("Check the country record.", "Confirm the code, name, language, timezone, and active status."),
        ("Add a local church.", "Under Local churches, add at least one active church for the readiness view."),
        ("Create the invitation.", "Under Administrators, choose Invite administrator, enter the person's name and email, and create the link."),
        ("Send the link securely.", "Copy the link and send it directly by email, WhatsApp, or another agreed channel. The link expires after seven days."),
    ]
    for title, detail in steps:
        add_number(doc, f"{title} {detail}", bold_prefix=f"{title} ")

    doc.add_heading("For the invited country administrator", level=2)
    invite_steps = [
        ("Open the invitation link", "in a modern browser."),
        ("Use the invited email address", "to create an account or sign in. A different email cannot accept the invitation."),
        ("Verify the email", "using the Firebase Authentication message."),
        ("Return to the invitation", "and accept access before the seven-day expiry."),
        ("Open the pilot app", "select the assigned country, choose the administration icon, and sign in."),
        ("Confirm the dashboard", "shows Country readiness, Country settings, Family Pledge, Holy Grounds, Songs, Payments, and Weekly videos."),
    ]
    for title, detail in invite_steps:
        add_number(doc, f"{title} {detail}", bold_prefix=f"{title} ")
    add_callout(
        doc,
        "Account safety",
        "Use one named account per administrator. Do not share passwords or reuse one national account among several people. Removing a person from one country does not delete that person's other country access.",
        fill=CREAM,
        border=GOLD,
    )

    doc.add_heading("7. Set up your country", level=1)
    add_para(doc, "Work through the sections below in order. Save each section, return to the public area, and confirm the result before continuing.")

    doc.add_heading("7.1 Country settings", level=2)
    for text in [
        "Open Administration > Country settings.",
        "Confirm the country name, main language, IANA timezone, and enabled status.",
        "Save, then return to the home page. Confirm the URL contains the country code, such as /de or /fr.",
        "Use the globe icon to test changing countries and the translate icon to test interface languages.",
    ]:
        add_number(doc, text)
    add_callout(
        doc,
        "Publishing effect",
        "A disabled country does not appear in the country selector. In this pilot there is no separate staging environment: saved, enabled content can become public for that country immediately.",
        fill="FCECEB",
        border=ERROR,
    )

    doc.add_heading("7.2 Family Pledge", level=2)
    for text in [
        "Open Administration > Family Pledge and choose Prepare default languages.",
        "Confirm Korean and English are present, then open the country's main language.",
        "Enter the approved title and all eight complete points.",
        "Turn on Language visible and save.",
        "Open the public Family Pledge and verify the language order, text, navigation, projection readability, and all eight points.",
    ]:
        add_number(doc, text)
    add_para(doc, "Readiness expects the country's main language, English, and Korean. If the main language is English or Korean, the required set naturally contains only the distinct languages.", italic=True, color=MUTED, size=10)

    doc.add_heading("7.3 Songs and audio", level=2)
    for text in [
        "Open Administration > Songs and select Standard catalog to prepare the bundled Worship starting catalog.",
        "For the first bulk load, download XLSX template or select Import XLSX/PPTX. The initial bulk import can be confirmed only once per country.",
        "Review every detected title, page, category, language, verse, and warning before Confirm import.",
        "Use New song for later additions. Enter title, page (leave blank if none), language, sort order, category, verses, chorus mode, and active status.",
        "Add audio individually. Each active track needs a name and an HTTPS URL or valid bundled asset path; verify verse-change timing where used.",
        "Use the active switch to hide a song without deleting its data, then test search, lyrics, playback, and offline availability in the public catalog.",
    ]:
        add_number(doc, text)
    add_callout(
        doc,
        "Import limits",
        "Use one song per PPTX. PPTX imports text only, not design, animation, images, fonts, or audio. A page value of 0 is not accepted in new edits; leave the field empty if there is no page.",
    )

    doc.add_heading("7.4 Offerings and tithes", level=2)
    for text in [
        "Open Administration > Payments and tithes.",
        "Enter the public page title, introduction, and optional final note.",
        "Add one or more methods, such as bank transfer, PIX, MB Way, a payment link, or Other.",
        "For each active method, add a clear label and at least one complete detail, payment URL, or QR content value.",
        "Enable the page, save, and test the public view. Copy each value, open each link, and scan each QR code with a separate device.",
    ]:
        add_number(doc, text)
    add_para(doc, "Only publish financial details that have been approved by the responsible national finance contact.", bold=True, color=ERROR)

    doc.add_heading("7.5 Weekly videos", level=2)
    for text in [
        "Open Administration > Weekly videos.",
        "Expand Update video links, then paste one valid YouTube watch link and one valid Vimeo link.",
        "Save and play both public cards. Check title, thumbnail/embed behavior, sound, and full-screen behavior.",
    ]:
        add_number(doc, text)
    add_para(doc, "Weekly videos require internet and are updated publicly as soon as the links are saved.", italic=True, color=MUTED, size=10)

    doc.add_heading("7.6 Holy Grounds", level=2)
    for text in [
        "Open Administration > Holy Grounds and choose New place.",
        "Enter name, city, full address, summary, story, visit instructions, responsible person, and contact email.",
        "Upload a JPG, PNG, or WebP photograph of no more than 5 MB.",
        "If adding coordinates, enter latitude and longitude together and verify the map destination.",
        "Turn on Place visible in directory, save, then search and filter for the entry in the public European directory.",
    ]:
        add_number(doc, text)
    add_para(doc, "Holy Grounds are optional for country launch readiness. Their text is not translated automatically, and replacing a Cloudinary image does not automatically delete the old file.", italic=True, color=MUTED, size=10)

    doc.add_heading("8. A quick check before inviting more people", level=1)
    add_para(doc, "These two screens are helpful while you are preparing your country. Think of them as reminders of what still needs attention:")
    add_table(
        doc,
        ["View", "Audience", "Checks"],
        [
            ["Country readiness", "Country admin", "Country active; admin access; songs; Pledge; payments; weekly videos; offline audio; Holy Grounds (optional)."],
            ["Operational CRM readiness", "Superadmin", "Country active; assigned admin; active local church; Pledge; remote song catalog; payments; videos; Holy Grounds (optional)."],
        ],
        [2200, 1750, 5410],
        font_size=9.2,
    )
    add_bullet(doc, "A green or complete indicator checks structure and format; it cannot tell whether an external video, audio link, payment instruction, or translation is right for your community.")
    add_bullet(doc, "Open every video, audio URL, QR value, image, translation, and address yourself before sharing it more widely.")
    add_bullet(doc, "If an optional section is not ready yet, simply leave it for later and make a note of who will add it.")

    doc.add_heading("9. Things to try as a member", level=1)
    add_para(doc, "Open an incognito/private browser window if you can. It gives you a fresh first-visit experience and helps you notice anything stored from earlier admin use.")
    add_table(
        doc,
        ["Try this", "What you should see", "Notes for your team"],
        [
            ["First visit and country selection", "Enabled country appears; selection opens the correct /country-code URL.", ""],
            ["Interface language", "Country language loads by default; translate icon switches all supported interface languages.", ""],
            ["Country separation", "Public content belongs to the selected country; a country admin cannot save another country's data.", ""],
            ["Family Pledge", "Correct visible languages, title, eight points, and navigation.", ""],
            ["Songs", "Search, category/page metadata, lyrics, chorus sequence, playback, timing, and active/hidden state are correct.", ""],
            ["Offerings", "Approved methods display; copy, links, and QR scanning work.", ""],
            ["Weekly videos", "YouTube and Vimeo cards load and play while online.", ""],
            ["Holy Grounds", "Search/filter works; current country sorts first; details and Google Maps destination are correct.", ""],
            ["Phone and computer", "Text and controls are easy to use without being cut off.", ""],
            ["Sign-out", "Admin sign-out removes editing access and public browsing remains available.", ""],
        ],
        [2200, 5040, 2120],
        font_size=8.7,
    )

    doc.add_heading("10. Sunday Mode: use it in a real service", level=1)
    doc.add_heading("Before the service", level=2)
    for text in [
        "Open the app online and select the correct country.",
        "Choose Start Sunday Mode.",
        "Enable only the modules used in that service and drag them into the local order.",
        "Wait for the offline-audio panel to report that audio is available offline; use Retry if needed.",
        "Open each module once during preparation and test full-screen presentation on the actual projection equipment.",
    ]:
        add_number(doc, text)
    doc.add_heading("During and after the service", level=2)
    for text in [
        "Start the service and open each enabled module from Sunday Mode.",
        "Returning from a module marks it complete; correct the checkbox manually if required.",
        "Choose Finish service and record whether the service ran without paper, whether it was uninterrupted, a 1-5 rating, and concise notes.",
        "If possible, use the same browser/device for a few real Sundays. The local summary becomes more useful as it collects repeated feedback.",
    ]:
        add_number(doc, text)
    add_callout(
        doc,
        "Local data",
        "The Sunday plan and reports stay on the browser/device used for the service. They do not synchronize to the CRM. Share a short note with the pilot coordinator if you discover something your team would like improved.",
        fill=CREAM,
        border=GOLD,
    )

    doc.add_heading("11. Using FFPMU without a connection", level=1)
    for text in [
        "While online, open the selected country and wait for current content to synchronize.",
        "Open the songs catalog and allow all enabled audio tracks to finish preparing. Visit Holy Ground images that must be available later.",
        "Turn off Wi-Fi and mobile data, or use the browser's offline mode.",
        "Reload the app and verify the last synchronized Family Pledge, songs, payments, Holy Grounds, and prepared audio.",
        "Confirm that weekly videos clearly remain an online-only function.",
        "Reconnect, reopen the app, and confirm new changes synchronize without losing the local Sunday report history.",
    ]:
        add_number(doc, text)
    add_para(doc, "Offline behavior depends on a successful earlier synchronization. A first-time visitor cannot expect content that has never been downloaded.", bold=True, color=DARK_GREEN)

    doc.add_heading("12. How to share feedback or report a problem", level=1)
    add_para(doc, "A helpful message gives us enough information to understand what happened. Include these details when you can:")
    for item in [
        "Country and role used (public user, country admin, or superadmin).",
        "Date/time, device, operating system, browser, and connection state.",
        "Exact page or module and the steps that reproduce the problem.",
        "What you expected and what actually happened.",
        "Screenshot or short screen recording, with private or financial data hidden.",
        "Frequency: always, sometimes, or once.",
        "Impact: blocks Sunday service; major workaround; minor problem; or suggestion.",
    ]:
        add_bullet(doc, item)
    add_callout(doc, "Suggested subject", "[FFPMU] COUNTRY - PAGE OR MODULE - short description")

    doc.add_heading("13. Quick help", level=1)
    add_table(
        doc,
        ["Problem", "Check first"],
        [
            ["Country is not listed", "Confirm the document code matches the country code and Country enabled is on; reload while online."],
            ["Access denied", "Use the invited, verified email; confirm the invitation was accepted and the correct country is assigned."],
            ["Invitation cannot be accepted", "Check the seven-day expiry, verified email, and exact email match; ask the superadmin for a new invitation if needed."],
            ["Pledge is not ready", "Confirm the main language, English, and Korean each have a title, eight complete points, and visible status."],
            ["Song import is unavailable", "The initial bulk import may already be complete; add or edit songs individually to avoid duplicates."],
            ["Audio still needs internet", "Open the direct HTTPS URL in the browser, wait for preparation, and use Retry. Invalid or blocked URLs cannot be cached."],
            ["Payment page is not ready", "Every active method needs a title plus at least one detail, link, or QR content value."],
            ["Holy Ground will not publish", "Add a valid image and required fields; enter both coordinates together and keep them within valid ranges."],
            ["Video does not play", "Use valid YouTube/Vimeo watch URLs and confirm the content is externally available and embeddable."],
        ],
        [2900, 6460],
        font_size=9.0,
    )

    doc.add_heading("14. A simple first week", level=1)
    add_para(doc, "There is no paperwork to complete. A comfortable way to begin is:")
    for item in [
        "Day 1: choose your country, look around the public pages, and make a note of content that belongs to your country.",
        "Day 2: accept the admin invitation and add the basic country settings and Family Pledge languages.",
        "Day 3: add a small set of songs, then check the public catalog on a phone and a computer.",
        "Before Sunday: add any approved offerings and video links your community wants to use; leave optional Holy Grounds for later if needed.",
        "Sunday: try Sunday Mode with the modules your local service actually uses. Afterwards, send one short message about what felt easy and what could be better.",
    ]:
        add_bullet(doc, item)
    add_callout(
        doc,
        "Remember",
        "Each country controls only its own public content. Take your time, keep financial information accurate, and ask the pilot coordinator whenever something is unclear.",
        fill=CREAM,
        border=GOLD,
    )
    add_link_line(doc, "Open FFPMU", APP_URL)
    add_link_line(doc, "Project repository", GITHUB_URL, "access may need to be granted")

    doc.core_properties.title = "FFPMU Country User and Admin Guide"
    doc.core_properties.subject = "Practical guidance for country setup, administration, and Sunday use"
    doc.core_properties.author = "FFPMU Pilot Team"
    doc.core_properties.keywords = "FFPMU, country admin, Sunday Mode, pilot, Europe"
    doc.save(GUIDE_PATH)


def build_community():
    doc = Document()
    configure_styles(doc)
    add_running_header(doc.sections[0], "FFPMU | Youth community", "Draft post")
    add_page_number_footer(doc.sections[0], "FFPMU Youth Community Post")

    add_title_block(
        doc,
        "Paste-ready community message",
        "FFPMU Youth Community Post",
        "A short, informal introduction for the European youth community.",
        [
            ("Audience", "European youth community"),
            ("Tone", "Friendly and informal"),
            ("Format", "One short post"),
            ("Prepared", DATE_LABEL),
        ],
        title_size=25,
    )
    add_callout(
        doc,
        "Before posting",
        "The app link is live. The GitHub link currently returns a 404 to people who are not signed in with access, so make the repository public before sharing this message.",
        fill="FCECEB",
        border=ERROR,
    )

    doc.add_heading("Message to share", level=1)
    add_para(doc, "Hi everyone!", bold=True, color=DARK_GREEN, size=12, after=10)
    add_para(
        doc,
        "After seeing the Family Pledge app shared here recently, I thought it was time to share something I have been building too 😄",
    )
    add_para(
        doc,
        "It is called FFPMU Connect, a web app to help with Sunday services. It has Sunday Mode, Family Pledge, the motto, songs with lyrics and audio, offerings, weekly videos, and a European Holy Grounds directory.",
    )
    add_para(
        doc,
        "The cool part is that every country has its own space and its own admin. Local leaders can manage their language, Pledge, songs, payment information, videos, and Holy Grounds without changing another country's content. Portugal is the first active test country, and I would love to try it with youth and leaders from other European countries.",
    )
    add_para(
        doc,
        "It is still a web pilot, so please tell me if something is wrong or if you have ideas. Help with translations, local songs, content, and testing it during a real Sunday service would be amazing! Want to test the admin side for your country? Use Request administrator access in the app and send me your country, name, role, and email.",
    )
    add_para(
        doc,
        "And before you open the GitHub 😂 I am an Engineering Manager now, so I do not code as much as I used to. Please do not judge the code too much hahaha. Feedback, ideas, and pull requests are very welcome!",
    )
    add_link_line(doc, "app", APP_URL)
    add_link_line(doc, "github", GITHUB_URL)
    add_para(doc, "And thanks to the member who shared the separate Family Pledge app too: ", after=0)
    p = doc.paragraphs[-1]
    add_hyperlink(p, FAMILY_PLEDGE_URL, FAMILY_PLEDGE_URL)
    add_para(doc, "Thank you!", bold=True, color=DARK_GREEN, before=8, after=3)
    add_para(doc, "Osman", italic=True, color=MUTED, after=0)

    doc.core_properties.title = "FFPMU Youth Community Post"
    doc.core_properties.subject = "Short paste-ready community message"
    doc.core_properties.author = "FFPMU Pilot Team"
    doc.core_properties.keywords = "FFPMU, youth, Europe, community, country admin"
    doc.save(COMMUNITY_PATH)


if __name__ == "__main__":
    build_guide()
    build_community()
    print(GUIDE_PATH)
    print(COMMUNITY_PATH)
