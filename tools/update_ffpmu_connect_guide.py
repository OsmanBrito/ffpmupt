from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


SOURCE = Path('deliverables/FFPMU_Country_Leader_Pilot_Training_Guide.docx')
OUTPUT = Path('deliverables/FFPMU_Connect_Country_Leader_Pilot_Training_Guide.docx')
COUNTRY_SELECTION_SCREENSHOT = Path(
    'docs/ffpmu_guide_screenshots/01-country-selection-connect.jpg'
)


def replace_paragraph(document, old, new):
    for paragraph in document.paragraphs:
        if paragraph.text == old:
            paragraph.text = new
            return paragraph
    raise ValueError(f'Paragraph not found: {old!r}')


def insert_after(anchor, text, style):
    element = OxmlElement('w:p')
    anchor._p.addnext(element)
    paragraph = Paragraph(element, anchor._parent)
    paragraph.style = style
    paragraph.add_run(text)
    return paragraph


def replace_table_text(document, old, new):
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text == old:
                    cell.text = new
                    return
    raise ValueError(f'Table cell not found: {old!r}')


def replace_first_screenshot(document, image_path):
    """Keep the existing layout while replacing the outdated branded capture."""
    first_shape = document.inline_shapes[0]
    width = first_shape.width
    height = first_shape.height
    screenshot_paragraph = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph._p.xpath('.//w:drawing')
    )
    screenshot_paragraph.clear()
    screenshot_paragraph.add_run().add_picture(
        str(image_path),
        width=width,
        height=height,
    )


def restart_numbered_list(document, template, paragraphs):
    """Give the inserted motto steps their own 1–4 Word numbering sequence."""
    template_num_pr = template._p.pPr.numPr
    template_num_id = str(template_num_pr.numId.val)
    numbering = document.part.numbering_part.element

    numbers = numbering.findall(qn('w:num'))
    new_num_id = max(int(number.get(qn('w:numId'))) for number in numbers) + 1
    template_number = next(
        number
        for number in numbers
        if number.get(qn('w:numId')) == template_num_id
    )
    new_number = deepcopy(template_number)
    new_number.set(qn('w:numId'), str(new_num_id))

    level_override = OxmlElement('w:lvlOverride')
    level_override.set(qn('w:ilvl'), '0')
    start_override = OxmlElement('w:startOverride')
    start_override.set(qn('w:val'), '1')
    level_override.append(start_override)
    new_number.append(level_override)
    numbering.append(new_number)

    for paragraph in paragraphs:
        paragraph_properties = paragraph._p.get_or_add_pPr()
        existing_numbering = paragraph_properties.find(qn('w:numPr'))
        if existing_numbering is not None:
            paragraph_properties.remove(existing_numbering)
        number_properties = OxmlElement('w:numPr')
        level = OxmlElement('w:ilvl')
        level.set(qn('w:val'), '0')
        number_id = OxmlElement('w:numId')
        number_id.set(qn('w:val'), str(new_num_id))
        number_properties.extend((level, number_id))
        paragraph_properties.append(number_properties)


document = Document(SOURCE)
replace_first_screenshot(document, COUNTRY_SELECTION_SCREENSHOT)

# Brand the guide with the new app name while retaining FFPMU for the
# organisation/community references that are not the product name.
replacements = {
    'FFPMU Country\nUser & Admin Guide': 'FFPMU Connect\nCountry User & Admin Guide',
    'A friendly, step-by-step guide to setting up your country and using FFPMU together.':
        'A friendly, step-by-step guide to setting up your country and using FFPMU Connect together.',
    '1. What FFPMU is': '1. What FFPMU Connect is',
    "FFPMU is a web-based Sunday service guide and country content hub. Members choose a country and receive that country's language and public content. The same shared application can support different national communities without mixing their administrative data.":
        "FFPMU Connect is a web-based Sunday service guide and country content hub. Members choose a country and receive that country's language and public content. The same shared application can support different national communities without mixing their administrative data.",
    'Open the pilot app select the assigned country, choose the administration icon, and sign in.':
        'Open the FFPMU Connect pilot app, select the assigned country, choose the administration icon, and sign in.',
    'Confirm the dashboard shows Country readiness, Country settings, Family Pledge, Holy Grounds, Songs, Payments, and Weekly videos.':
        'Confirm the dashboard shows Country readiness, Country settings, Family Pledge, Motto, Holy Grounds, Songs, Payments, and Weekly videos.',
    '11. Using FFPMU without a connection': '11. Using FFPMU Connect without a connection',
    'Suggested subject: [FFPMU] COUNTRY - PAGE OR MODULE - short description':
        'Suggested subject: [FFPMU Connect] COUNTRY - PAGE OR MODULE - short description',
    'Open FFPMU: https://ffpmupt-402e1.web.app/':
        'Open FFPMU Connect: https://ffpmupt-402e1.web.app/',
}

for old, new in replacements.items():
    replace_paragraph(document, old, new)

# Make the access-request path visible before any admin terminology appears.
# Leaders do not need Firebase access or a pre-created account: the project
# coordinator sends an invitation after receiving the basic country details.
pilot_link = next(
    paragraph
    for paragraph in document.paragraphs
    if paragraph.text.startswith('Pilot app: https://ffpmupt-402e1.web.app/')
)
insert_after(
    pilot_link,
    'Need administrator access? Open Request administrator access on the first screen, select your country, and submit your full name, role (country leader/admin), and the individual email address you will use. The form opens a prepared email; press Send in your email app. You do not need Firebase access or a shared password.',
    'Callout',
)

replace_paragraph(
    document,
    'Country content is independent. Editing or disabling a song, Pledge language, payment method, or Holy Ground in one country does not change another country.',
    'Country content is independent. Editing a song, Pledge language, motto, payment method, or Holy Ground in one country does not change another country.',
)

replace_paragraph(
    document,
    'Yearly motto, country-specific offerings/tithes, and weekly YouTube/Vimeo videos.',
    'A country-specific yearly motto, offerings/tithes, and weekly YouTube/Vimeo videos.',
)

replace_paragraph(
    document,
    'The sample countries in this image are for illustration; your app will show the countries that are currently enabled.',
    'The countries in this image are for illustration; your app will show the countries that are currently enabled.',
)

# Add the new country-admin motto workflow and keep the remaining section
# numbers consecutive.
for old, new in {
    '7.3 Songs and audio': '7.4 Songs and audio',
    '7.4 Offerings and tithes': '7.5 Offerings and tithes',
    '7.5 Weekly videos': '7.6 Weekly videos',
    '7.6 Holy Grounds': '7.7 Holy Grounds',
}.items():
    replace_paragraph(document, old, new)

anchor = next(
    paragraph
    for paragraph in document.paragraphs
    if paragraph.text == 'Readiness expects the country\'s main language, English, and Korean. If the main language is English or Korean, the required set naturally contains only the distinct languages.'
)
heading_source = next(
    paragraph for paragraph in document.paragraphs if paragraph.text == '7.4 Songs and audio'
)
body_source = next(
    paragraph
    for paragraph in document.paragraphs
    if paragraph.text == 'Open Administration > Songs and select Standard catalog to prepare the bundled Worship starting catalog.'
)

heading = insert_after(anchor, '7.3 Yearly motto', heading_source.style)
step_one = insert_after(heading, 'Open Administration > Motto.', body_source.style)
step_two = insert_after(
    step_one,
    'Enter the title and the message shown on the public motto page. Keep the wording approved for your country.',
    body_source.style,
)
step_three = insert_after(
    step_two,
    'Save, then open the public Motto page or Sunday Mode and check that the updated title and text are visible.',
    body_source.style,
)
step_four = insert_after(
    step_three,
    'Country boundary: this change affects only the selected country. Until a country saves its own motto, the current default motto is shown.',
    body_source.style,
)

restart_numbered_list(document, body_source, [step_one, step_two, step_three, step_four])

# Add an explicit request path to the access section as well as the opening
# callout, so a leader can find it again while reading the admin workflow.
access_heading = next(
    paragraph
    for paragraph in document.paragraphs
    if paragraph.text == '6. Administrator invitation and first sign-in'
)
superadmin_heading = next(
    paragraph
    for paragraph in document.paragraphs
    if paragraph.text == 'For the superadministrator'
)
request_heading = insert_after(
    access_heading,
    'If you need access',
    superadmin_heading.style,
)
request_body = insert_after(
    request_heading,
    'Open Request administrator access from the first screen, select your country, and submit your full name, role, and individual email. The form opens a prepared email to the coordinator; press Send, then the coordinator reviews it and creates the official invitation. Then follow the steps under “For the invited country administrator”.',
    body_source.style,
)
insert_after(
    request_body,
    'Important: Do not create a Firebase account or share an administrator password unless I have first sent you an invitation.',
    'Callout',
)

replace_table_text(
    document,
    'Manage settings, readiness, Pledge, songs, payments, videos, and Holy Grounds.',
    'Manage settings, readiness, Pledge, motto, songs, payments, videos, and Holy Grounds.',
)

document.core_properties.title = 'FFPMU Connect Country User & Admin Guide'
document.core_properties.subject = 'Pilot training guide for FFPMU Connect country leaders and administrators'
document.core_properties.comments = 'Updated for FFPMU Connect branding and country-specific motto management.'
document.save(OUTPUT)

print(OUTPUT)
